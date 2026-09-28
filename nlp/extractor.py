"""
extractor.py
------------
Extracts text from PDF or DOCX resumes and records
structural signals that can affect ATS compatibility.
"""

from dataclasses import dataclass, field
from pathlib import Path
import re


@dataclass
class ExtractionResult:
    text: str = ""
    file_type: str = ""
    used_tables: bool = False
    used_text_boxes: bool = False
    used_headers_footers: bool = False
    likely_scanned_image: bool = False
    multi_column_layout: bool = False
    warnings: list = field(default_factory=list)


def extract(file_path: str) -> ExtractionResult:

    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf(path)

    elif suffix == ".docx":
        return _extract_docx(path)

    else:
        raise ValueError(
            f"Unsupported file type: {suffix}. "
            "Use .pdf or .docx"
        )


def _extract_pdf(path: Path) -> ExtractionResult:

    import pdfplumber

    result = ExtractionResult(
        file_type="pdf"
    )

    full_text_parts = []
    x0_positions = []

    with pdfplumber.open(path) as pdf:

        for page in pdf.pages:

            page_text = page.extract_text() or ""

            full_text_parts.append(page_text)

            # Detect tables
            if page.find_tables():
                result.used_tables = True

            # Collect word positions
            words = page.extract_words()

            for word in words:
                x0_positions.append(
                    round(word["x0"])
                )

    result.text = "\n".join(
        full_text_parts
    ).strip()

    # Detect scanned/image-only PDF
    if not result.text:
        result.likely_scanned_image = True

    # Basic multi-column heuristic
    if x0_positions:

        distinct_starts = sorted(
            set(x0_positions)
        )

        if len(distinct_starts) > 1:

            gaps = [
                distinct_starts[i + 1] -
                distinct_starts[i]
                for i in range(
                    len(distinct_starts) - 1
                )
            ]

            if gaps and max(gaps) > 150:
                result.multi_column_layout = True

    return result


def _extract_docx(path: Path) -> ExtractionResult:

    import docx

    result = ExtractionResult(
        file_type="docx"
    )

    document = docx.Document(
        str(path)
    )

    body_text = []

    # Normal paragraphs
    for para in document.paragraphs:

        body_text.append(
            para.text
        )

    # Tables
    if document.tables:

        result.used_tables = True

        for table in document.tables:

            for row in table.rows:

                for cell in row.cells:

                    body_text.append(
                        cell.text
                    )

    # Headers and footers
    for section in document.sections:

        header_text = "\n".join(
            p.text
            for p in section.header.paragraphs
            if p.text.strip()
        )

        footer_text = "\n".join(
            p.text
            for p in section.footer.paragraphs
            if p.text.strip()
        )

        if (
            header_text.strip()
            or footer_text.strip()
        ):
            result.used_headers_footers = True

    # Text boxes
    xml = document.element.xml

    if (
        "<w:txbxContent>" in xml
        or re.search(
            r"<wps:txbx>",
            xml
        )
    ):
        result.used_text_boxes = True

    # Final extracted text
    result.text = "\n".join(
        text
        for text in body_text
        if text.strip()
    )

    # No normal text warning
    if not result.text.strip():

        result.warnings.append(
            "Little to no plain paragraph text was found. "
            "Check whether content is trapped in text boxes, "
            "images, or SmartArt."
        )

    return result